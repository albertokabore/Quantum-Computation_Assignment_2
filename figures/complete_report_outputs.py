"""Complete the existing corrected report with reproducible code and outputs."""
import contextlib
import io
import json
from pathlib import Path
import matplotlib
matplotlib.use('Agg')
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

root = Path(__file__).resolve().parent.parent
path = root / 'Assignment2_Report_Final.docx'
doc = Document(path)
nb = json.loads((root / 'Assignment2_Quantum_Circuits.ipynb').read_text(encoding='utf-8'))
scope = {}
outputs = {}
for i, cell in enumerate(nb['cells']):
    if cell['cell_type'] == 'code':
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            exec(compile(''.join(cell['source']), f'notebook-cell-{i}', 'exec'), scope)
        outputs[i] = out.getvalue()

def after(p, text, style=None):
    node = OxmlElement('w:p')
    p._p.addnext(node)
    q = Paragraph(node, p._parent)
    q.text = text
    if style:
        q.style = style
    return q

def replace_range(start, end, code):
    paragraphs = doc.paragraphs
    first = paragraphs[start]
    style = first.style
    first.text = code
    for p in paragraphs[start+1:end+1]:
        p._p.getparent().remove(p._p)
    return first, style

# Replace the old appendix with the actual notebook implementation.
paragraphs = doc.paragraphs
start = next(i for i,p in enumerate(paragraphs) if p.text == 'import os' and i > 200)
end = next(i for i,p in enumerate(paragraphs) if p.text.startswith('qc_z, state_z'))
imports = 'import os\nimport numpy as np\nimport matplotlib.pyplot as plt\nfrom IPython.display import display\nfrom qiskit import QuantumCircuit, transpile\nfrom qiskit.quantum_info import Statevector\nfrom qiskit.visualization import plot_histogram\nfrom qiskit_aer import AerSimulator\nos.makedirs("figures", exist_ok=True)\nSHOTS, SEED = 1024, 42\n'
appendix = imports + ''.join(nb['cells'][5]['source']) + '\n' + '\n'.join(''.join(nb['cells'][i]['source']) for i in (7,9,11))
p, style = replace_range(start, end, appendix)
out = io.StringIO()
with contextlib.redirect_stdout(out):
    exec(compile(appendix, 'report-appendix', 'exec'), {})
q = after(p, 'Output of Appendix A (the circuit diagrams and histograms are also included in Sections 3.3–3.5):')
after(q, out.getvalue().strip(), style)

# Complete each individual gate example with explicit simulation output.
for gate in 'hyz':
    paragraphs = doc.paragraphs
    gate_index = next(i for i,p in enumerate(paragraphs) if p.text == f'qc.{gate}(0)')
    start, end = gate_index - 1, gate_index + 2
    code = f'''qc = QuantumCircuit(1, 1)
qc.{gate}(0)
state = Statevector.from_instruction(qc)
qc.measure(0, 0)
print(qc.draw())
sim = AerSimulator()
raw = sim.run(transpile(qc, sim, seed_transpiler=42), shots=1024, seed_simulator=42).result().get_counts()
counts = {{k: raw.get(k, 0) for k in ("0", "1")}}
print("Statevector:", np.round(state.data, 4))
print("Counts:", counts)
print("Observed frequencies:", {{k: v / 1024 for k, v in counts.items()}})
display(plot_histogram(counts))'''
    p, style = replace_range(start, end, code)
    out = io.StringIO()
    with contextlib.redirect_stdout(out):
        exec(compile(code, f'report-{gate}', 'exec'), scope)
    q = after(p, 'Output (the generated histogram appears below with its mathematical verification):')
    after(q, out.getvalue().strip(), style)

# Pair the overview example with its exact output.
paragraphs = doc.paragraphs
start = next(i for i,p in enumerate(paragraphs) if p.text == 'from qiskit import QuantumCircuit, transpile')
end = next(i for i,p in enumerate(paragraphs[start:], start) if p.text == 'plot_histogram(counts)')
code = imports + '''qc = QuantumCircuit(1, 1)
qc.h(0)
state = Statevector.from_instruction(qc)
qc.measure(0, 0)
print(qc.draw())
sim = AerSimulator()
counts = sim.run(transpile(qc, sim, seed_transpiler=42), shots=1024, seed_simulator=42).result().get_counts()
print("Statevector:", np.round(state.data, 4))
print("Counts:", counts)
display(plot_histogram(counts))'''
p, style = replace_range(start, end, code)
out = io.StringIO()
with contextlib.redirect_stdout(out):
    exec(compile(code, 'report-overview', 'exec'), {})
q = after(p, 'Output (the H histogram is reproduced in Figure 1b):')
after(q, out.getvalue().strip(), style)

# Remove the obsolete duplicate normalization snippet with no matching output.
for p in list(doc.paragraphs):
    if p.text in ('a1 = np.sqrt(3) * np.exp(1j * np.pi / 6) / 2', 'psi = Statevector([a0, a1])', 'print(abs(a0)**2, abs(a1)**2, abs(a0)**2 + abs(a1)**2, psi.is_valid())'):
        p._p.getparent().remove(p._p)
    elif p.text.startswith('a0 = 1j / 2'):
        p.text = 'import numpy as np\nfrom qiskit.quantum_info import Statevector\n\n' + p.text
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            exec(compile(p.text, 'report-normalization', 'exec'), {})
        assert 'Qiskit Statevector.is_valid(): True' in out.getvalue()

doc.save(path)
verified = Document(path)
assert len(verified.inline_shapes) == 6
assert len(verified._element.xpath('.//m:oMath')) > 20
assert not any('screenshot' in p.text.lower() for p in verified.paragraphs)
assert sum(p.text.startswith('Output') for p in verified.paragraphs) == 6
print('Updated existing report:', path.name)
print('Six runnable code examples verified with matching output; six figures and equations preserved.')
