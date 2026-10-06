"""Expand only the requested assignment explanations and embed real screenshots."""
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Inches
from PIL import Image

root = Path(__file__).resolve().parent.parent
path = root / 'Assignment2_Report_Revised.docx'
doc = Document(path)

def after(prefix, texts):
    p = next(p for p in doc.paragraphs if p.text.startswith(prefix))
    for text in texts:
        node = OxmlElement('w:p')
        p._p.addnext(node)
        p = Paragraph(node, p._parent)
        p.text = text

after('Question 1 -', [
    'The question gives |ψ⟩ = (i/2)|0⟩ + (√3 e^(iπ/6)/2)|1⟩. Use the computational basis |0⟩ = (1, 0)ᵀ and |1⟩ = (0, 1)ᵀ. These vectors are orthonormal: ⟨0|0⟩ = ⟨1|1⟩ = 1 and ⟨0|1⟩ = ⟨1|0⟩ = 0.',
    'The conjugate transpose is ⟨ψ| = (−i/2)⟨0| + (√3 e^(−iπ/6)/2)⟨1|. Expanding ⟨ψ|ψ⟩ gives |a₀|²⟨0|0⟩ + a₀* a₁⟨0|1⟩ + a₁* a₀⟨1|0⟩ + |a₁|²⟨1|1⟩. The two cross terms vanish because the basis vectors are orthogonal.',
    '|a₀|² = (−i/2)(i/2) = −i²/4 = 1/4. Also, |a₁|² = (√3/2)² e^(−iπ/6)e^(iπ/6) = (3/4)e⁰ = 3/4. Thus ⟨ψ|ψ⟩ = 1/4 + 3/4 = 1 and ‖ψ‖ = √1 = 1. This proves the required normalization without treating complex amplitudes as ordinary real squares.',
    'As an independent algebra check, e^(iπ/6) = √3/2 + i/2, so a₁ = 3/4 + i√3/4. Its squared magnitude is (3/4)² + (√3/4)² = 9/16 + 3/16 = 3/4.'
])
after('Question 2 -', [
    'The three requested inputs are |0⟩, independently of the state in Question 1. H is the Hadamard gate; Y and Z are Pauli gates. Although the instructions recall X, they request implementations only of H|0⟩, Y|0⟩ and Z|0⟩.',
    'For an output b₀|0⟩ + b₁|1⟩, computational-basis measurement obeys the Born rule: P(|0⟩) = |b₀|² and P(|1⟩) = |b₁|². Derive these amplitudes by matrix multiplication before comparing them with simulator output. State vectors are computed before measurement; the measured circuit then produces classical counts.'
])
after('(a) H|0', [
    'Detailed hand calculation: H = (1/√2)[[1, 1], [1, −1]]. Multiplying by (1, 0)ᵀ gives (1/√2)(1·1 + 1·0, 1·1 + (−1)·0)ᵀ = (1/√2, 1/√2)ᵀ. Hence H|0⟩ = (|0⟩ + |1⟩)/√2.',
    'Both amplitudes are 1/√2, so P(|0⟩) = (1/√2)*(1/√2) = 1/2 and P(|1⟩) = 1/2. Their sum is 1. For 1024 shots the expected counts are 512 and 512; expected counts are averages over repeated experiments, not a requirement that every run gives exactly these counts.'
])
after('(b) Y|0', [
    'Detailed hand calculation: Y = [[0, −i], [i, 0]]. Therefore Y(1, 0)ᵀ = (0·1 + (−i)·0, i·1 + 0·0)ᵀ = (0, i)ᵀ = i|1⟩.',
    'The output amplitudes are b₀ = 0 and b₁ = i. Thus P(|0⟩) = 0 and P(|1⟩) = i* i = (−i)i = 1. The factor i multiplies the entire state by a global phase, so it does not change the outcome probabilities. Every ideal shot yields 1.'
])
after('(c) Z|0', [
    'Detailed hand calculation: Z = [[1, 0], [0, −1]]. Therefore Z(1, 0)ᵀ = (1·1 + 0·0, 0·1 + (−1)·0)ᵀ = (1, 0)ᵀ = |0⟩.',
    'The output amplitudes are b₀ = 1 and b₁ = 0, so P(|0⟩) = |1|² = 1 and P(|1⟩) = |0|² = 0. Z changes the sign of a |1⟩ component, but the requested input has no such component. Every ideal shot yields 0.'
])
after('Comparison of predictions', [
    'The H run gives frequencies 526/1024 = 0.513671875 and 498/1024 = 0.486328125. Their deviations from 0.5 are +0.013671875 and −0.013671875. For independent shots with p = 0.5, the standard deviation of the count of zeros is √(1024 × 0.5 × 0.5) = 16. The observed difference of 14 counts from 512 is 0.875 standard deviations, consistent with finite sampling. This comparison supports the prediction; the exact state-vector check independently verifies the amplitudes.',
    'Y gives 0/1024 = 0 for outcome 0 and 1024/1024 = 1 for outcome 1. Z gives 1024/1024 = 1 for outcome 0 and 0/1024 = 0 for outcome 1. Both match the predictions exactly. These are ideal AerSimulator results, not measurements from a physical quantum processor.'
])

for shape, (gate, kind) in zip(doc.inline_shapes, [(g,k) for g in 'hyz' for k in ('circuit','histogram')]):
    image_path = root / 'figures' / f'screenshot_{kind}_{gate}.png'
    rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    doc.part.related_parts[rid]._blob = image_path.read_bytes()
    width, height = Image.open(image_path).size
    shape.width = Inches(3.4 if kind == 'circuit' else 5.8)
    shape.height = int(shape.width * height / width)

for p in doc.paragraphs:
    if p.text.startswith('Figure ') and 'Qiskit circuit diagram' in p.text:
        p.text = p.text.replace('Qiskit circuit diagram', 'Screenshot of the executed notebook circuit output')
    elif p.text.startswith('Figure ') and 'Simulated measurement histogram' in p.text:
        p.text = p.text.replace('Simulated measurement histogram', 'Screenshot of the executed notebook measurement histogram')

# Remove stale claims about a bonus X run: the assignment requests H, Y and Z.
for table in doc.tables:
    for row in list(table.rows):
        if '(bonus)' in ' '.join(c.text for c in row.cells):
            table._tbl.remove(row._tr)
    for row in table.rows:
        for cell in row.cells:
            if 'bonus X circuit' in cell.text:
                cell.text = ('Why does the factor i not show up in the histogram?\n'
                             'Y|0⟩ = i|1⟩ differs from |1⟩ by a global phase. '
                             'The squared magnitude is |i|² = (−i)i = 1, so '
                             'computational-basis measurement gives outcome 1 with certainty.')
            elif 'Measured P(0), P(1)' == cell.text:
                cell.text = 'Observed f(0), f(1)'

assert len(doc.inline_shapes) == 6
doc.save(path)
print('Expanded report saved:', path)
