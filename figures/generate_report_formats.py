"""Build the editable Word report while preserving the final PDF's scanned work."""
from copy import deepcopy
from pathlib import Path
import argparse
import io

import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
REVIEW = ROOT / "figures" / "report_export_review"
SOURCE_PDF = REVIEW / "source_report.pdf"
REPORT_PDF = ROOT / "Kabore_Albert_Module_2_Assignment.pdf"
REPORT_WORD = ROOT / "Kabore_Albert_Module_2_Assignment.docx"
NAVY = "203B60"


def page_images(pdf, page_index):
    return [b for b in pdf[page_index].get_text("dict")["blocks"] if b["type"] == 1]


def make_word(screenshot=None):
    pdf = pymupdf.open(SOURCE_PDF)
    template = Document(ROOT / "Assignment2_Report_Revised.docx")
    doc = Document()
    doc.core_properties.title = "Quantum Computation — Module 2 Assignment"
    doc.core_properties.author = "Albert Kabore"
    doc.core_properties.subject = "Single-qubit states and H, Y, Z gates in Qiskit"
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.65)
    section.header_distance = section.footer_distance = Inches(0.3)
    normal = doc.styles["Normal"]
    normal.font.name = "Cambria"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in (("Title", 27), ("Heading 1", 15), ("Heading 2", 12)):
        style = doc.styles[name]
        style.font.name = "Cambria"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(NAVY)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(10)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Quantum Computation · Assignment 2  |  Page ")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 100, 100)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)

    def text(value, style=None):
        return doc.add_paragraph(value, style)

    def heading(value, level=1, new_page=False):
        p = text(value, f"Heading {level}")
        p.paragraph_format.page_break_before = new_page
        return p

    def picture(data, width=6.5, height_limit=None, new_page=False):
        im = Image.open(io.BytesIO(data))
        if height_limit:
            width = min(width, height_limit * im.width / im.height)
        p = text("")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.page_break_before = new_page
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = 1
        p.add_run().add_picture(io.BytesIO(data), width=Inches(width))
        return p

    def caption(value):
        p = text(value)
        p.paragraph_format.space_after = Pt(10)
        for run in p.runs:
            run.italic = True
            run.font.size = Pt(9)
        return p

    def scanned_page(index):
        # Only existing handwritten pages are rasterized; report text stays editable.
        data = pdf[index].get_pixmap(matrix=pymupdf.Matrix(2, 2)).tobytes("png")
        picture(data, width=6.75, height_limit=9.5, new_page=True)

    def copy_table(index, widths):
        node = deepcopy(template.tables[index]._tbl)
        doc._body._body.insert(len(doc._body._body) - 1, node)
        table = Table(node, doc._body)
        table.autofit = False
        for col, width in zip(table.columns, widths):
            col.width = Inches(width)
        for row in table.rows:
            no_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(no_split)
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.05
                    for run in p.runs:
                        run.font.name = "Cambria"
                        run.font.size = Pt(9)
        return table

    # Page 1: retain the report's author and original assignment date.
    title = text("Quantum Computation\nModule 2 Assignment", "Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(130)
    p = text("Single-Qubit State Vectors, Pauli Gates and the Hadamard Gate in Qiskit")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = text("Albert Kabore")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p = text("October 5, 2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(65)

    # Page 2: the supplied normalization proof.
    scanned_page(1)

    # Page 3: assignment statement and editable gate-action table.
    heading("Question 2", new_page=True)
    text("Implement (a) H|0⟩, (b) Y|0⟩, and (c) Z|0⟩ using qc = QuantumCircuit(1, 1). Each circuit below follows the assignment order:")
    text("(i) Measure the output and show the histogram screenshot.\n(ii) Show the circuit screenshot produced by draw().\n(iii) Derive the output state and P(|0⟩), P(|1⟩), then compare with the histogram.\n(iv) Collect all answers in this report, alongside the supplied Question 1 answer.")
    text("The computational basis is |0⟩ = (1, 0)ᵀ and |1⟩ = (0, 1)ᵀ. X, Y, and Z are the Pauli operators; H is the Hadamard gate. The imaginary unit satisfies i² = −1.")
    text("X = [[0, 1], [1, 0]]\nY = [[0, −i], [i, 0]]\nZ = [[1, 0], [0, −1]]")
    heading("Gate actions and Qiskit syntax", 2)
    copy_table(4, [0.6, 1.1, 1.55, 1.55, 2.4])

    # Page 4: reproduce the shared runnable code from the existing final PDF.
    heading("Qiskit implementation", new_page=True)
    text("Each circuit starts in |0⟩ and has one classical bit. The output state vector is computed before measurement. Measurement in the computational basis stores the result in classical bit 0. All reported runs use an ideal AerSimulator with 1024 shots and seed 42.")
    text("For b₀|0⟩ + b₁|1⟩, the Born rule gives P(|0⟩) = |b₀|² and P(|1⟩) = |b₁|². Observed frequencies are counts divided by 1024.")
    code = '''import numpy as np
from IPython.display import display
from qiskit import QuantumCircuit, transpile
from qiskit.quantum_info import Statevector
from qiskit.visualization import plot_histogram
from qiskit_aer import AerSimulator

def run_circuit(gate):
    qc = QuantumCircuit(1, 1)
    getattr(qc, gate)(0)
    state = Statevector.from_instruction(qc)
    qc.measure(0, 0)
    sim = AerSimulator()
    compiled = transpile(qc, sim, seed_transpiler=42)
    raw = sim.run(compiled, shots=1024,
                  seed_simulator=42).result().get_counts()
    counts = {k: raw.get(k, 0) for k in ("0", "1")}
    print("Statevector:", np.round(state.data, 4))
    print("Counts:", counts)
    print("Observed frequencies:",
          {k: v / 1024 for k, v in counts.items()})
    display(plot_histogram(counts))
    display(qc.draw("mpl"))
    return qc, state, counts'''
    p = text(code)
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_together = True
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), "F4F6F8")
    p._p.get_or_add_pPr().append(shade)
    for run in p.runs:
        run.font.name = "Consolas"
        run.font.size = Pt(9)

    # Pages 5–13: histogram, circuit, handwritten derivation for each gate.
    gates = [
        ("h", "(a)", "Hadamard", 4, 5, 6, "{'0': 526, '1': 498}", "[0.7071+0.j, 0.7071+0.j]"),
        ("y", "(b)", "Pauli Y", 7, 8, 9, "{'0': 0, '1': 1024}", "[0.+0.j, 0.+1.j]"),
        ("z", "(c)", "Pauli Z", 10, 11, 12, "{'0': 1024, '1': 0}", "[1.+0.j, 0.+0.j]"),
    ]
    for gate, part, name, hist_page, circ_page, deriv_page, counts, state in gates:
        ket = f"{gate.upper()}|0⟩"
        heading(f"{part} {ket} — {name} gate applied to |0⟩", new_page=True)
        heading(f"(i) Measurement histogram for {ket}", 2)
        text(f'Apply qc.{gate}(0) to QuantumCircuit(1, 1). Run qc_{gate}, state_{gate}, counts_{gate} = run_circuit("{gate}") using the shared setup. Measure with qc.measure(0, 0).')
        p = text(f"1024-shot output: {counts}\nState vector before measurement: {state}")
        p.paragraph_format.space_after = Pt(10)
        hist_images = page_images(pdf, hist_page)
        picture(hist_images[0]["image"], width=5.25, height_limit=2.05)
        shot = screenshot.read_bytes() if gate == "h" and screenshot else hist_images[1]["image"]
        picture(shot, width=6.6, height_limit=3.85)
        caption(f"{ket}: measurement histogram and executed-notebook screenshot.")

        heading(f"(ii) Circuit screenshot for {ket} after running draw()", new_page=True)
        text(f'The measured circuit is drawn with qc.draw("mpl"). It contains one {gate.upper()} gate on qubit 0 followed by measurement into classical bit 0.')
        circ_images = page_images(pdf, circ_page)
        picture(circ_images[0]["image"], width=3.4, height_limit=1.8)
        shot = screenshot.read_bytes() if gate == "h" and screenshot else circ_images[1]["image"]
        picture(shot, width=6.6, height_limit=4.8)
        caption(f"{ket}: circuit diagram and executed-notebook screenshot.")
        scanned_page(deriv_page)

    # Page 14: editable native table, including native Word math from the template.
    heading("(iv) One document containing all answers", new_page=True)
    text("This report contains the supplied Question 1 answer and the answers to (a), (b), and (c), including code, histogram and circuit screenshots, mathematical predictions, and verification. The table summarizes the results for 1024 shots per circuit.")
    copy_table(7, [0.65, 1.65, 1.05, 1.4, 1.15, 1.3])
    text("")
    text("H agrees within finite-shot variation; Y and Z agree exactly with the ideal predictions. These are ideal AerSimulator results, not measurements from a physical quantum processor.")
    heading("Reference", 2)
    text("IBM Quantum Learning, Basics of Quantum Information — Single systems.\nhttps://learning.quantum.ibm.com/course/basics-of-quantum-information/single-systems")
    doc.save(REPORT_WORD)
    pdf.close()
    check = Document(REPORT_WORD)
    assert len(check.inline_shapes) == 16
    assert len(check.tables) == 2
    assert len(check._element.xpath(".//m:oMath")) >= 15
    print(f"Word report created: {REPORT_WORD.name}")
    print("Editable report text and tables; 16 preserved illustrations; native table equations.")


def replace_pdf_screenshot(screenshot):
    pdf = pymupdf.open(SOURCE_PDF)
    data = screenshot.read_bytes()
    for index, area in (
        (4, pymupdf.Rect(80.9, 378, 531.1, 708)),
        (5, pymupdf.Rect(48, 322, 564, 701)),
    ):
        page = pdf[index]
        old = page_images(pdf, index)[1]
        # Remove only this image occurrence; retain all page text and the first diagram.
        page.add_redact_annot(pymupdf.Rect(old["bbox"]), fill=(1, 1, 1))
        page.apply_redactions(images=2, graphics=0, text=1)
        page.insert_image(area, stream=data, keep_proportion=True)
    temp = REVIEW / "report_with_replacement_screenshot.pdf"
    pdf.save(temp, garbage=4, deflate=True)
    pdf.close()
    with pymupdf.open(temp) as check:
        assert len(check) == 14
        assert "526" in check[4].get_text()
        assert "1024" in check[13].get_text()
        check[4].get_pixmap(matrix=pymupdf.Matrix(1.2, 1.2)).save(REVIEW / "updated_h_histogram_page.png")
        check[5].get_pixmap(matrix=pymupdf.Matrix(1.2, 1.2)).save(REVIEW / "updated_h_circuit_page.png")
    temp.replace(REPORT_PDF)
    print(f"Both Hadamard notebook screenshots replaced in {REPORT_PDF.name} (14 pages).")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--screenshot", type=Path)
    parser.add_argument("--update-pdf", action="store_true")
    args = parser.parse_args()
    if args.screenshot and not args.screenshot.is_file():
        parser.error("The supplied screenshot path does not exist.")
    if args.update_pdf and not args.screenshot:
        parser.error("A screenshot is required before updating the PDF.")
    make_word(args.screenshot)
    if args.update_pdf:
        replace_pdf_screenshot(args.screenshot)
