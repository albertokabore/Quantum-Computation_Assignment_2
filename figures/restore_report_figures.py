"""Restore generated figures and align the report with its companion notebook."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches
from PIL import Image

root = Path(__file__).resolve().parent.parent
path = root / 'Assignment2_Report_Reviewed.docx'
notebook_path = root / 'Assignment2_Quantum_Circuits.ipynb'
notebook = json.loads(notebook_path.read_text(encoding='utf-8'))
doc = Document(path)
assert len(doc.inline_shapes) == 6
for shape, (gate, kind) in zip(doc.inline_shapes, [(g,k) for g in 'hyz' for k in ('circuit','histogram')]):
    image_path = root / 'figures' / f'{kind}_{gate}.png'
    rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    doc.part.related_parts[rid]._blob = image_path.read_bytes()
    width, height = Image.open(image_path).size
    shape.width = Inches(3.2 if kind == 'circuit' else 5.8)
    shape.height = int(shape.width * height / width)

for p in doc.paragraphs:
    if p.text.startswith('Figure '):
        text = p.text.replace('Screenshot of the executed notebook circuit output', 'Qiskit circuit diagram').replace('Screenshot of the executed notebook histogram', 'Simulated measurement histogram')
        p.text = text
    elif 'The screenshots show counts' in p.text:
        p.text = p.text.replace('The screenshots show counts', 'The histograms show counts')
    elif p.text.startswith('import os\n') and 'def run_single_qubit_circuit' in p.text:
        p.text = '\n'.join([
            'import os\nimport numpy as np\nimport matplotlib.pyplot as plt\nfrom IPython.display import display\nfrom qiskit import QuantumCircuit, transpile\nfrom qiskit.quantum_info import Statevector\nfrom qiskit.visualization import plot_histogram\nfrom qiskit_aer import AerSimulator',
            'os.makedirs("figures", exist_ok=True)\nSHOTS, SEED = 1024, 42',
            ''.join(notebook['cells'][5]['source']),
            '\n'.join(''.join(notebook['cells'][i]['source']) for i in (7,9,11)),
        ])
    elif p.text.startswith('a1 = np.sqrt(3)') and 'print(abs(a0)' in p.text:
        p._element.getparent().remove(p._element)

try:
    doc.save(path)
except PermissionError:
    path = root / 'Assignment2_Report_Final.docx'
    doc.save(path)
last = ''.join(notebook['cells'][-1]['source'])
start = last.index('The Word report embeds')
end = last.index('\n\nReferences:', start)
last = last[:start] + 'The Word report includes three generated Qiskit circuit diagrams and three simulated measurement histograms for H, Y, and Z, together with the code, mathematical derivations, probabilities, and verification. Figures are exported directly from Qiskit/Matplotlib; no screenshots are included.' + last[end:]
last = last.replace('Assignment2_Report_Reviewed.docx', path.name)
notebook['cells'][-1]['source'] = last.splitlines(keepends=True)
notebook_path.write_text(json.dumps(notebook, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
verified = Document(path)
assert len(verified.inline_shapes) == 6
assert not any('screenshot' in p.text.lower() for p in verified.paragraphs)
assert len(verified._element.xpath('.//m:oMath')) > 20
print('Restored six generated figures; screenshot references removed; equations preserved.')
print('Saved:', path)
